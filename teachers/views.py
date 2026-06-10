import io
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib import messages
from django.http import HttpResponse
from django.utils import timezone
from django.db.models import Count, Q
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from exams.models import ExamResult, StudentAnswer, ExamSubject


def teacher_login(request):
    if request.user.is_authenticated and request.user.is_staff:
        return redirect('teacher_dashboard')

    if request.method == 'POST':
        username = request.POST.get('username', '').strip()
        password = request.POST.get('password', '')
        user = authenticate(request, username=username, password=password)
        if user and user.is_staff:
            login(request, user)
            return redirect('teacher_dashboard')
        messages.error(request, 'Неверный логин/пароль или недостаточно прав')

    return render(request, 'teachers/login.html')


def teacher_logout(request):
    logout(request)
    return redirect('teacher_login')


@staff_member_required(login_url='/teacher/login/')
def teacher_dashboard(request):
    sort = request.GET.get('sort', 'date')

    sort_map = {
        'date':    '-start_time',
        'student': 'student__last_name',
        'exam':    'exam__name',
        'score':   '-score',
        'pending': '-pending_count',
    }

    results = list(
        ExamResult.objects.exclude(status='in_progress')
        .select_related('student', 'exam', 'exam__course')
        .prefetch_related('student_answers__question')
        .annotate(
            pending_count=Count(
                'student_answers',
                filter=Q(
                    student_answers__is_correct__isnull=True,
                    student_answers__question__question_type__in=['open', 'text']
                )
            )
        )
        .order_by(sort_map.get(sort, '-start_time'))
    )

    diff_labels = {'easy': 'Лёгкие', 'medium': 'Средние', 'hard': 'Сложные'}
    for result in results:
        stats = {k: {'correct': 0, 'total': 0, 'points': 0, 'label': v}
                 for k, v in diff_labels.items()}
        for sa in result.student_answers.all():
            d = sa.question.difficulty
            if d in stats:
                stats[d]['total'] += 1
                if sa.is_correct:
                    stats[d]['correct'] += 1
                stats[d]['points'] += sa.effective_points
        result.difficulty_stats = {k: v for k, v in stats.items() if v['total'] > 0}

    return render(request, 'teachers/dashboard.html', {
        'results': results,
        'current_sort': sort,
        'sort_options': {
            'date':    'По дате',
            'score':   'По баллам',
            'student': 'По студенту',
            'exam':    'По экзамену',
            'pending': 'Непроверенные первые',
        },
    })


@staff_member_required(login_url='/teacher/login/')
def export_results_excel(request):
    results = list(
        ExamResult.objects.exclude(status='in_progress')
        .select_related('student', 'exam', 'exam__course')
        .prefetch_related('student_answers__question')
        .annotate(
            pending_count=Count(
                'student_answers',
                filter=Q(
                    student_answers__is_correct__isnull=True,
                    student_answers__question__question_type__in=['open', 'text']
                )
            )
        )
        .order_by('-score')
    )

    diff_labels = {'easy': 'Лёгкие', 'medium': 'Средние', 'hard': 'Сложные'}
    for result in results:
        stats = {k: {'correct': 0, 'total': 0, 'points': 0} for k in diff_labels}
        for sa in result.student_answers.all():
            d = sa.question.difficulty
            if d in stats:
                stats[d]['total'] += 1
                if sa.is_correct:
                    stats[d]['correct'] += 1
                stats[d]['points'] += sa.effective_points
        result.diff_stats = stats

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Результаты экзаменов'

    # Стили
    header_font = Font(bold=True, color='FFFFFF', size=11)
    header_fill = PatternFill('solid', fgColor='2D6A9F')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left = Alignment(horizontal='left', vertical='center')
    thin = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    diff_fills = {
        'easy':   PatternFill('solid', fgColor='D4EDDA'),
        'medium': PatternFill('solid', fgColor='FFF3CD'),
        'hard':   PatternFill('solid', fgColor='F8D7DA'),
    }

    # Заголовки
    headers = [
        ('#', 4),
        ('ID студента', 14),
        ('Студент', 24),
        ('Экзамен', 28),
        ('Курс', 18),
        ('Дата сдачи', 16),
        ('Балл', 8),
        ('Макс. балл', 10),
        ('% выполнения', 13),
        ('Лёгкие (верн/всего)', 20),
        ('Средние (верн/всего)', 21),
        ('Сложные (верн/всего)', 21),
        ('На проверке', 13),
        ('Статус', 14),
    ]

    ws.row_dimensions[1].height = 30
    for col, (title, width) in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.freeze_panes = 'A2'

    # Данные
    status_map = {'finished': 'Завершён', 'time_expired': 'Время вышло'}
    for i, result in enumerate(results, start=1):
        pct = round(result.score / result.max_score * 100) if result.max_score else 0
        easy = result.diff_stats['easy']
        medium = result.diff_stats['medium']
        hard = result.diff_stats['hard']

        row_data = [
            i,
            result.student.student_id,
            result.student.full_name,
            result.exam.name,
            result.exam.course.name,
            result.start_time.strftime('%d.%m.%Y %H:%M') if result.start_time else '',
            result.score,
            result.max_score,
            pct,
            f"{easy['correct']}/{easy['total']}" if easy['total'] else '—',
            f"{medium['correct']}/{medium['total']}" if medium['total'] else '—',
            f"{hard['correct']}/{hard['total']}" if hard['total'] else '—',
            result.pending_count,
            status_map.get(result.status, result.status),
        ]

        row_num = i + 1
        ws.row_dimensions[row_num].height = 18
        for col, value in enumerate(row_data, start=1):
            cell = ws.cell(row=row_num, column=col, value=value)
            cell.border = border
            cell.alignment = center if col not in (2, 3, 4, 5) else left

        # Цвет % выполнения
        pct_cell = ws.cell(row=row_num, column=9)
        if pct >= 80:
            pct_cell.fill = diff_fills['easy']
        elif pct >= 50:
            pct_cell.fill = diff_fills['medium']
        else:
            pct_cell.fill = diff_fills['hard']

        # Чередование фона строк
        if i % 2 == 0:
            row_fill = PatternFill('solid', fgColor='F5F8FC')
            for col in range(1, len(headers) + 1):
                cell = ws.cell(row=row_num, column=col)
                if not cell.fill or cell.fill.fgColor.rgb in ('00000000', 'FFFFFFFF'):
                    cell.fill = row_fill

    # Итоговая строка
    last_row = len(results) + 2
    ws.cell(row=last_row, column=1, value='Итого').font = Font(bold=True)
    ws.cell(row=last_row, column=1).alignment = center
    ws.cell(row=last_row, column=7, value=sum(r.score for r in results)).font = Font(bold=True)
    ws.cell(row=last_row, column=7).alignment = center

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    filename = f"results_{timezone.now().strftime('%Y%m%d_%H%M')}.xlsx"
    response = HttpResponse(
        buf.read(),
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = f'attachment; filename="{filename}"'
    return response


@staff_member_required(login_url='/teacher/login/')
def teacher_review(request, exam_result_id):
    exam_result = get_object_or_404(ExamResult, pk=exam_result_id)

    open_answers = exam_result.student_answers.filter(
        question__question_type__in=['open', 'text']
    ).select_related('question', 'question__subject', 'reviewed_by')

    if request.method == 'POST':
        for sa in open_answers:
            score_key = f'score_{sa.id}'
            comment_key = f'comment_{sa.id}'
            score_val = request.POST.get(score_key, '').strip()
            comment_val = request.POST.get(comment_key, '').strip()

            if score_val != '':
                try:
                    sa.teacher_score = float(score_val)
                except ValueError:
                    pass
            else:
                sa.teacher_score = None

            sa.teacher_comment = comment_val
            sa.reviewed_at = timezone.now()
            if hasattr(request.user, 'teacher_profile'):
                sa.reviewed_by = request.user.teacher_profile
            sa.save()

        exam_result.recalculate_score()
        messages.success(request, 'Оценки сохранены, итоговый балл пересчитан.')
        return redirect('teacher_review', exam_result_id=exam_result_id)

    for sa in open_answers:
        es = ExamSubject.objects.filter(
            exam=exam_result.exam, subject=sa.question.subject
        ).first()
        if es:
            if sa.question.difficulty == 'easy':
                sa.max_points = es.easy_points
            elif sa.question.difficulty == 'medium':
                sa.max_points = es.medium_points
            else:
                sa.max_points = es.hard_points
        else:
            sa.max_points = 0

    return render(request, 'teachers/review.html', {
        'exam_result': exam_result,
        'open_answers': open_answers,
    })


@staff_member_required(login_url='/teacher/login/')
def teacher_result_detail(request, exam_result_id):
    """Полный отчёт студента, доступный преподавателю."""
    exam_result = get_object_or_404(ExamResult, pk=exam_result_id)

    student_answers = exam_result.student_answers.select_related(
        'question', 'question__subject'
    ).prefetch_related('question__answers', 'selected_answers')

    subject_stats = {}
    difficulty_stats = {
        'easy':   {'correct': 0, 'total': 0, 'points': 0, 'max_points': 0, 'label': 'Лёгкие'},
        'medium': {'correct': 0, 'total': 0, 'points': 0, 'max_points': 0, 'label': 'Средние'},
        'hard':   {'correct': 0, 'total': 0, 'points': 0, 'max_points': 0, 'label': 'Сложные'},
    }

    for answer in student_answers:
        subject = answer.question.subject.name if answer.question.subject else 'Без предмета'
        if subject not in subject_stats:
            subject_stats[subject] = {'correct': 0, 'total': 0, 'points': 0}
        subject_stats[subject]['total'] += 1
        if answer.is_correct:
            subject_stats[subject]['correct'] += 1
        subject_stats[subject]['points'] += answer.effective_points

        diff = answer.question.difficulty
        if diff in difficulty_stats:
            difficulty_stats[diff]['total'] += 1
            if answer.is_correct:
                difficulty_stats[diff]['correct'] += 1
            difficulty_stats[diff]['points'] += answer.effective_points

    return render(request, 'exams/exam_result_detail.html', {
        'exam_result': exam_result,
        'student_answers': student_answers,
        'subject_stats': subject_stats,
        'difficulty_stats': difficulty_stats,
        'is_teacher_view': True,
    })
